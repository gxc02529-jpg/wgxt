# -*- coding: utf-8 -*-
"""文稿批量导入服务：解析 Excel（封面/标题/栏目分类/正文）并逐行校验导入草稿

规则：
- 标题为空 / 正文为空 → 该行失败；标题与库内重复 → 跳过。
- 栏目按名称匹配（支持顿号/逗号/斜杠分隔多栏目），匹配不到的忽略并提示。
- 封面为空或下载失败 → 封面留空，文稿仍导入（不因封面问题丢弃整篇）。
- 导入结果统一为草稿（draft），作者为当前用户。
"""
import io
import os
import re

import httpx
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..config import settings
from ..models.article import Article
from ..models.category import Category
from ..models.user import User
from ..services.oss_service import get_oss_config
from ..services.storage_service import save_bytes
from ..services.text_utils import strip_html

# 表头别名 → 标准列名
HEADER_MAP = {
    "封面": "cover", "封面图": "cover", "图片": "cover",
    "标题": "title", "题目": "title",
    "分类": "category", "栏目": "category", "栏目分类": "category",
    "正文": "content", "内容": "content", "正文内容": "content",
}

ALLOWED_IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp", "image/svg+xml"}
CONTENT_TYPE_EXT = {
    "image/jpeg": ".jpg", "image/jpg": ".jpg",
    "image/png": ".png", "image/gif": ".gif",
    "image/webp": ".webp", "image/svg+xml": ".svg",
}

MAX_FILE_SIZE = 5 * 1024 * 1024   # Excel 文件上限 5MB
MAX_IMPORT_ROWS = 500             # 单次最多导入 500 行，防止滥用
MAX_COVER_SIZE = settings.MAX_UPLOAD_SIZE  # 封面图大小上限（与上传一致 20MB）


def _parse_header(ws) -> dict:
    """识别表头：返回 {标准列名: 列索引}，找不到的列不包含"""
    cols = {}
    for row in ws.iter_rows(values_only=True, max_row=1):
        for idx, cell in enumerate(row or []):
            if cell is None:
                continue
            key = HEADER_MAP.get(str(cell).strip())
            if key and key not in cols:
                cols[key] = idx
    return cols


def _match_categories(db: Session, text: str) -> tuple[list, list]:
    """按名称匹配栏目（支持顿号/逗号/斜杠/分号分隔），返回 (匹配到的栏目列表, 未匹配的名称列表)"""
    matched, missed = [], []
    for name in re.split(r"[、,，/／;；]", text):
        name = name.strip()
        if not name:
            continue
        cat = db.query(Category).filter(
            Category.name == name, Category.status == "active"
        ).first()
        if cat:
            matched.append(cat)
        else:
            missed.append(name)
    return matched, missed


def _looks_like_html(text: str) -> bool:
    """粗略判断文本是否为 HTML：出现 <p>/<div>/<br> 等标签时按 HTML 处理"""
    return bool(re.search(r"<[a-zA-Z/][^>]*>", text))


def _download_cover(oss_cfg, url: str) -> str:
    """下载封面图并按存储配置保存，返回可访问 URL；失败抛异常由调用方记录"""
    if not url.startswith(("http://", "https://")):
        raise ValueError("仅支持 http/https 链接")
    with httpx.Client(follow_redirects=True, timeout=15) as client:
        with client.stream("GET", url) as resp:
            if resp.status_code != 200:
                raise ValueError(f"HTTP {resp.status_code}")
            content_type = (resp.headers.get("content-type") or "").split(";")[0].strip().lower()
            if content_type not in ALLOWED_IMAGE_TYPES:
                # 部分站点不返回 content-type，按 URL 后缀兜底
                ext_from_url = os.path.splitext(url.split("?")[0])[1].lower()
                ext_to_type = {
                    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
                    ".gif": "image/gif", ".webp": "image/webp", ".svg": "image/svg+xml",
                }
                content_type = ext_to_type.get(ext_from_url, "")
            if content_type not in ALLOWED_IMAGE_TYPES:
                raise ValueError("非支持的图片格式")
            ext = CONTENT_TYPE_EXT.get(content_type, ".jpg")
            buf = io.BytesIO()
            total = 0
            for chunk in resp.iter_bytes(64 * 1024):
                total += len(chunk)
                if total > MAX_COVER_SIZE:
                    raise ValueError(f"图片过大，最大 {MAX_COVER_SIZE // 1024 // 1024}MB")
                buf.write(chunk)
            if total == 0:
                raise ValueError("图片内容为空")
    _, url_stored, _ = save_bytes(oss_cfg, buf.getvalue(), content_type, ext)
    return url_stored


def import_excel(db: Session, user: User, file_bytes: bytes, filename: str) -> dict:
    """解析并导入 Excel，返回统计结果：total/imported/skipped/failed/warnings"""
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "仅支持 .xlsx 格式的 Excel 文件")
    if len(file_bytes) > MAX_FILE_SIZE:
        raise HTTPException(400, "文件过大，最大允许 5MB")
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析 Excel 文件，请确认是有效的 .xlsx")

    ws = wb.active
    cols = _parse_header(ws)
    if "title" not in cols or "content" not in cols:
        wb.close()
        raise HTTPException(400, "Excel 表头需要包含「标题」和「正文」列")

    oss_cfg = get_oss_config(db)
    imported, skipped, failed, warnings = 0, [], [], []
    row_no = 0
    for row in ws.iter_rows(values_only=True):
        row_no += 1
        if row_no == 1:
            continue  # 表头
        if imported + len(skipped) + len(failed) >= MAX_IMPORT_ROWS:
            warnings.append({"row": row_no, "reason": f"超过单次导入上限 {MAX_IMPORT_ROWS} 行，其余行未处理"})
            break
        # 跳过全空行
        if row is None or all(c is None or not str(c).strip() for c in row):
            continue

        title = str(row[cols["title"]] or "").strip()
        content = str(row[cols["content"]] or "").strip()
        if not title:
            failed.append({"row": row_no, "reason": "标题为空"})
            continue
        if not content:
            failed.append({"row": row_no, "reason": "正文为空"})
            continue
        if db.query(Article).filter(Article.title == title).first():
            skipped.append({"row": row_no, "reason": "标题已存在"})
            continue

        # 栏目匹配（缺失列或匹配不到不阻断导入）
        cat_text = ""
        if cols.get("category") is not None and row[cols["category"]] is not None:
            cat_text = str(row[cols["category"]]).strip()
        cats, missed = _match_categories(db, cat_text) if cat_text else ([], [])
        if missed:
            warnings.append({"row": row_no, "reason": f"栏目未匹配：{'、'.join(missed)}"})

        # 封面：为空或下载失败都不丢弃文稿
        cover_url = ""
        if cols.get("cover") is not None and row[cols["cover"]] is not None:
            cover_raw = str(row[cols["cover"]]).strip()
            if cover_raw:
                try:
                    cover_url = _download_cover(oss_cfg, cover_raw)
                except Exception as e:
                    warnings.append({"row": row_no, "reason": f"封面下载失败，已留空：{str(e)[:50]}"})

        # 正文与摘要处理：兼容纯文本与 HTML 两种来源格式
        if _looks_like_html(content):
            # 已是 HTML 正文（如 news-batch 导出的 <p> 段落 + <img>）：原样入库，
            # 纯文本与摘要从 HTML 剥离生成，避免摘要/纯文本里残留标签
            html = content
            plain = strip_html(content)
            content_source = plain
            first_paras = re.findall(r"<p[^>]*>(.*?)</p>", content, re.S)
            if first_paras:
                summary = strip_html(first_paras[0])[:200]
            else:
                summary = plain[:200]
        else:
            # 纯文本正文：按换行转 <p> 段落；摘要取正文首段前 200 字
            paragraphs = [p.strip() for p in content.splitlines() if p.strip()]
            html = "".join(f"<p>{p}</p>" for p in paragraphs)
            summary = paragraphs[0][:200] if paragraphs else ""
            content_source = content

        article = Article(
            title=title,
            summary=summary,
            content=html,
            content_source=content_source,
            cover_url=cover_url,
            author_id=user.id,
            status="draft",
        )
        db.add(article)
        db.flush()
        article.categories = cats
        db.commit()
        imported += 1

    wb.close()
    return {
        "total": imported + len(skipped) + len(failed),
        "imported": imported,
        "skipped": skipped,
        "failed": failed,
        "warnings": warnings,
    }
